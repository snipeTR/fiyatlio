"""Rapor dosyaları ve CLI."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from fiyatlio.cli import main
from fiyatlio.report import write_reports
from fiyatlio.cli import analyze_files

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "samples" / "paxg_referans.csv"


def test_rapor_dosyalari(tmp_path):
    analysis = analyze_files(
        [SAMPLE],
        holdings={"PAXG": __import__("decimal").Decimal("0.671466")},
        prices={"PAXG": __import__("decimal").Decimal("4279.78")},
    )
    out = write_reports(analysis, tmp_path / "out")
    for name in (
        "ozet.md",
        "donemler.csv",
        "filller.csv",
        "kalan_lotlar.csv",
        "varlik_ozeti.csv",
        "rapor.html",
        "rapor.xlsx",
    ):
        assert (out / name).is_file(), name

    markdown = (out / "ozet.md").read_text(encoding="utf-8")
    assert "PAXG" in markdown
    assert "Açık pozisyon" in markdown
    assert "Kalan lotlar" in markdown
    assert "KISMI_SATIS" in (out / "donemler.csv").read_text(encoding="utf-8")

    html = (out / "rapor.html").read_text(encoding="utf-8")
    assert 'id="coin"' in html
    assert "PAXG/USDT" in html
    assert "function filtrele" in html

    book = load_workbook(out / "rapor.xlsx")
    assert "Dönemler" in book.sheetnames
    assert "Filller" in book.sheetnames
    assert "Kalan Lotlar" in book.sheetnames
    assert "Varlık Özeti" in book.sheetnames
    assert "Emirler" in book.sheetnames


def test_cli_analyze(tmp_path):
    out = tmp_path / "rapor"
    proc = subprocess.run(
        [
            sys.executable,
            "-m",
            "fiyatlio",
            "analyze",
            str(SAMPLE),
            "--holdings",
            "PAXG=0.671466",
            "--price",
            "PAXG=4279.78",
            "--method",
            "fifo",
            "--out",
            str(out),
        ],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert proc.returncode == 0, proc.stderr + proc.stdout
    assert "PAXG" in proc.stdout
    assert (out / "ozet.md").is_file()


def test_cli_hatali_holdings():
    try:
        main(["analyze", str(SAMPLE), "--holdings", "PAXG"])
    except SystemExit as exc:
        assert exc.code not in (0, None)
    else:
        raise AssertionError("hatalı cüzdan satırı kabul edildi")
